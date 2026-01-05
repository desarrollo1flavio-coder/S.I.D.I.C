"""
Exportador a Microsoft Word.

Genera documentos .docx con el informe completo.
"""
from pathlib import Path
from typing import Dict, Optional, Any
from datetime import date
import io

try:
    from docx import Document
    from docx.shared import Inches, Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

import pandas as pd

from ..models.report_data import ReportData


class WordExporter:
    """
    Exportador de reportes a formato Word (.docx).
    
    Genera un documento profesional con tablas y gráficos.
    """
    
    def __init__(self, report_data: ReportData):
        """
        Inicializa el exportador.
        
        Args:
            report_data: Datos del reporte a exportar.
        """
        if not HAS_DOCX:
            raise ImportError("python-docx no está instalado")
        
        self.report = report_data
        self.document = Document()
        self._last_error = ""
        self._setup_styles()
    
    def get_last_error(self) -> str:
        """Retorna el último error ocurrido."""
        return self._last_error
    
    def _setup_styles(self):
        """Configura los estilos del documento."""
        # Configurar márgenes
        sections = self.document.sections
        for section in sections:
            section.top_margin = Cm(2)
            section.bottom_margin = Cm(2)
            section.left_margin = Cm(2)
            section.right_margin = Cm(2)
    
    def _add_heading(self, text: str, level: int = 1):
        """Agrega un encabezado."""
        heading = self.document.add_heading(text, level=level)
        
        # Color rojo para títulos principales
        if level == 1:
            for run in heading.runs:
                run.font.color.rgb = RGBColor(204, 0, 0)
    
    def _add_paragraph(self, text: str, bold: bool = False, align: str = 'left'):
        """Agrega un párrafo."""
        p = self.document.add_paragraph()
        run = p.add_run(text)
        run.bold = bold
        
        if align == 'center':
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == 'right':
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        
        return p
    
    def _set_cell_shading(self, cell, color: str):
        """Aplica color de fondo a una celda."""
        shading = OxmlElement('w:shd')
        shading.set(qn('w:fill'), color)
        cell._tc.get_or_add_tcPr().append(shading)
    
    def _add_table(
        self,
        df: pd.DataFrame,
        title: Optional[str] = None
    ):
        """
        Agrega una tabla desde un DataFrame.
        
        Args:
            df: DataFrame con los datos.
            title: Título de la tabla (opcional).
        """
        if df.empty:
            return
        
        if title:
            self._add_heading(title, level=2)
        
        # Crear tabla
        n_rows = len(df) + 1  # +1 para encabezados
        n_cols = len(df.columns)
        
        table = self.document.add_table(rows=n_rows, cols=n_cols)
        table.style = 'Table Grid'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        # Encabezados
        header_row = table.rows[0]
        for col_idx, col_name in enumerate(df.columns):
            cell = header_row.cells[col_idx]
            cell.text = str(col_name)
            
            # Estilo de encabezado (rojo)
            self._set_cell_shading(cell, 'FF0000')
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.font.size = Pt(10)
        
        # Datos
        for row_idx, (_, row) in enumerate(df.iterrows()):
            table_row = table.rows[row_idx + 1]
            
            # Detectar si es fila de total
            first_val = str(row.iloc[0]).upper() if len(row) > 0 else ""
            is_total = 'TOTAL' in first_val and 'SUBTOTAL' not in first_val
            is_subtotal = 'SUBTOTAL' in first_val
            
            for col_idx, value in enumerate(row):
                cell = table_row.cells[col_idx]
                cell.text = str(value) if value is not None else ""
                
                # Aplicar estilos según tipo de fila
                if is_total:
                    self._set_cell_shading(cell, 'FFFF00')
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.bold = True
                elif is_subtotal:
                    self._set_cell_shading(cell, '0000FF')
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
                else:
                    # Alternar colores
                    color = 'FFFFFF' if row_idx % 2 == 0 else 'F0F0F0'
                    self._set_cell_shading(cell, color)
                
                # Centrar valores numéricos
                for paragraph in cell.paragraphs:
                    if col_idx > 0:  # Columnas numéricas
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    paragraph.paragraph_format.space_after = Pt(0)
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
        
        # Espacio después de la tabla
        self.document.add_paragraph()
    
    def _add_image(self, image_bytes: bytes, width: float = 5.5):
        """Agrega una imagen al documento."""
        try:
            self.document.add_picture(io.BytesIO(image_bytes), width=Inches(width))
            
            # Centrar la imagen
            last_paragraph = self.document.paragraphs[-1]
            last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        except Exception as e:
            print(f"Error insertando imagen: {e}")
    
    # ═══════════════════════════════════════════════════════════════════════
    # SECCIONES DEL DOCUMENTO
    # ═══════════════════════════════════════════════════════════════════════
    
    def _add_header(self):
        """Agrega encabezado del informe."""
        # Título principal
        self._add_heading(self.report.titulo, level=1)
        
        # Información del reporte
        periodo = self.report.periodo_principal
        
        if self.report.jurisdiccion:
            self._add_paragraph(f"Jurisdicción: {self.report.jurisdiccion}", bold=True)
        
        if periodo:
            self._add_paragraph(f"Período: {periodo.rango_fechas}", bold=True)
        
        self._add_paragraph(
            f"Fecha de generación: {date.today().strftime('%d/%m/%Y')}",
            bold=True
        )
        
        self.document.add_paragraph()  # Espacio
    
    def _add_resumen(self):
        """Agrega sección de resumen."""
        periodo = self.report.periodo_principal
        if not periodo:
            return
        
        self._add_heading("RESUMEN", level=2)
        
        resumen = [
            f"• Total de Hechos Delictuales: {periodo.total_hechos}",
            f"• Total de Personas Mencionadas: {periodo.total_mencionados}",
            f"• Total de Personas Aprehendidas: {periodo.total_aprehendidos}",
        ]
        
        # Categorías
        categorias = periodo.conteo_por_categoria()
        for cat, cant in categorias.items():
            resumen.append(f"  - {cat}: {cant}")
        
        for linea in resumen:
            self._add_paragraph(linea)
        
        self.document.add_paragraph()
    
    def _add_mencionados_texto(self):
        """Agrega sección de mencionados en formato texto."""
        periodo = self.report.periodo_principal
        if not periodo or not periodo.mencionados:
            return
        
        self._add_heading("MENCIONADOS COMO POSIBLES AUTORES MATERIALES", level=2)
        
        for m in periodo.mencionados:
            # Formato similar al de la imagen
            texto = (
                f"➤ {m.alias_formateado}, mencionado en un {m.delito}, "
                f"lugar del hecho: {m.direccion_hecho}, "
                f"el día {m.fecha.strftime('%d/%m/%Y') if m.fecha else 'S/F'} "
                f"a horas {m.hora.strftime('%H:%M') if m.hora else 'S/H'}."
            )
            
            if m.datos_filiatorios:
                texto += f" DATOS: {m.datos_filiatorios}"
            
            self._add_paragraph(texto)
        
        self.document.add_paragraph()
    
    # ═══════════════════════════════════════════════════════════════════════
    # EXPORTACIÓN PRINCIPAL
    # ═══════════════════════════════════════════════════════════════════════
    
    def export(
        self,
        output_path: str,
        tables: Dict[str, pd.DataFrame],
        charts: Optional[Dict[str, bytes]] = None,
        options: Optional[Dict[str, Any]] = None
    ) -> bool:
        """
        Exporta el reporte completo a Word.
        
        Args:
            output_path: Ruta del archivo de salida.
            tables: Diccionario de tablas generadas.
            charts: Diccionario de gráficos (bytes PNG).
            options: Opciones de exportación (incluir_graficos, incluir_cuadro_ref, etc.)
        
        Returns:
            True si se exportó correctamente.
        """
        charts = charts or {}
        options = options or {}
        
        # Opciones por defecto
        incluir_graficos = options.get('incluir_graficos', True)
        incluir_cuadro_ref = options.get('incluir_cuadro_ref', True)
        incluir_mencionados = options.get('incluir_mencionados', True)
        incluir_matrices = options.get('incluir_matrices', True)
        incluir_comparativos = options.get('incluir_comparativos', True)
        
        try:
            # Encabezado
            self._add_header()
            
            # Resumen
            self._add_resumen()
            
            # Cuadro de referencia
            if incluir_cuadro_ref and 'cuadro_referencia' in tables:
                df = tables['cuadro_referencia']
                # Simplificar para Word
                df_simple = df[['Símbolo', 'Descripción', 'Cantidad']].copy()
                self._add_table(df_simple, "CUADRO DE REFERENCIA")
            
            # Delitos con modalidades
            if 'delitos' in tables:
                self._add_table(tables['delitos'], "DELITOS CON MODALIDADES")
                if incluir_graficos and 'delitos' in charts:
                    self._add_image(charts['delitos'])
            
            # Días de la semana
            if 'dias_semana' in tables:
                self._add_table(tables['dias_semana'], "DÍAS DE LA SEMANA EN QUE OCURRIERON LOS HECHOS")
                if incluir_graficos and 'dias_semana' in charts:
                    self._add_image(charts['dias_semana'])
            
            # Franja horaria
            if 'franja_horaria' in tables:
                self._add_table(tables['franja_horaria'], "FRANJA HORARIA EN QUE OCURRIERON LOS HECHOS")
                if incluir_graficos and 'franja_horaria' in charts:
                    self._add_image(charts['franja_horaria'])
            
            # Movilidad
            if 'movilidad' in tables:
                self._add_table(tables['movilidad'], "MEDIOS DE MOVILIDAD UTILIZADOS")
                if incluir_graficos and 'movilidad' in charts:
                    self._add_image(charts['movilidad'])
            
            # Armas
            if 'armas' in tables:
                self._add_table(tables['armas'], "MEDIOS O ARMAS UTILIZADAS EN ROBOS AGRAVADOS")
                if incluir_graficos and 'armas' in charts:
                    self._add_image(charts['armas'])
            
            # Ámbito
            if 'ambito' in tables:
                self._add_table(tables['ambito'], "ÁMBITO DE OCURRENCIA DELICTUAL")
                if incluir_graficos and 'ambito' in charts:
                    self._add_image(charts['ambito'])
            
            # Matrices
            if incluir_matrices and 'matriz_delito_dia' in tables:
                self._add_table(tables['matriz_delito_dia'], "DELITOS POR DÍAS DE LA SEMANA")
                if incluir_graficos and 'delito_dia' in charts:
                    self._add_image(charts['delito_dia'])
            
            if incluir_matrices and 'matriz_delito_franja' in tables:
                self._add_table(tables['matriz_delito_franja'], "DELITOS POR FRANJA HORARIA")
                if incluir_graficos and 'delito_franja' in charts:
                    self._add_image(charts['delito_franja'])
            
            # Mencionados (formato texto)
            if incluir_mencionados:
                self._add_mencionados_texto()
            
            # Aprehendidos
            if 'aprehendidos' in tables:
                self._add_table(tables['aprehendidos'], "PERSONAS APREHENDIDAS")
            
            if 'aprehendidos_clasificacion' in tables:
                self._add_table(tables['aprehendidos_clasificacion'], "CLASIFICACIÓN DE APREHENDIDOS")
            
            # Comparativa
            if incluir_comparativos and 'comparativa_general' in tables and not tables['comparativa_general'].empty:
                self._add_table(tables['comparativa_general'], "CUADRO COMPARATIVO ENTRE PERÍODOS")
            
            # Guardar
            Path(output_path).parent.mkdir(parents=True, exist_ok=True)
            self.document.save(output_path)
            
            return True
        
        except Exception as e:
            import traceback
            import logging
            logger = logging.getLogger(__name__)
            error_msg = f"Error exportando a Word: {e}"
            logger.error(error_msg)
            traceback.print_exc()
            # Guardar el error para que pueda ser recuperado
            self._last_error = str(e)
            return False
