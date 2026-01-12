"""
Exportador a Microsoft Excel.

Genera archivos .xlsx con todas las tablas, gráficos y estilos
del informe delictual.
"""
import io
from pathlib import Path
from typing import Dict, List, Optional, Any
from datetime import date

try:
    from openpyxl import Workbook
    from openpyxl.styles import (
        Font, Fill, PatternFill, Border, Side, Alignment,
        NamedStyle
    )
    from openpyxl.utils.dataframe import dataframe_to_rows
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.chart import BarChart, Reference
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

import pandas as pd

from ..models.report_data import ReportData
from ..utils.constants import ColoresPolicial


class ExcelExporter:
    """
    Exportador de reportes a formato Excel con estilos profesionales.
    
    Genera un archivo .xlsx completo con todas las tablas y gráficos.
    """
    
    def __init__(self, report_data: ReportData):
        """
        Inicializa el exportador.
        
        Args:
            report_data: Datos del reporte a exportar.
        """
        if not HAS_OPENPYXL:
            raise ImportError("openpyxl no está instalado")
        
        self.report = report_data
        self.workbook = Workbook()
        self._setup_styles()
    
    def _setup_styles(self):
        """Configura los estilos de Excel."""
        # Estilo para encabezados (fondo rojo, texto blanco)
        self.header_style = NamedStyle(name='header')
        self.header_style.font = Font(bold=True, color='FFFFFF', size=11)
        self.header_style.fill = PatternFill(start_color='FF0000', end_color='FF0000', fill_type='solid')
        self.header_style.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        self.header_style.border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Estilo para totales (fondo amarillo)
        self.total_style = NamedStyle(name='total')
        self.total_style.font = Font(bold=True, color='000000', size=11)
        self.total_style.fill = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')
        self.total_style.alignment = Alignment(horizontal='center', vertical='center')
        self.total_style.border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Estilo para subtotales (fondo azul, texto blanco)
        self.subtotal_style = NamedStyle(name='subtotal')
        self.subtotal_style.font = Font(bold=True, color='FFFFFF', size=10)
        self.subtotal_style.fill = PatternFill(start_color='0000FF', end_color='0000FF', fill_type='solid')
        self.subtotal_style.alignment = Alignment(horizontal='center', vertical='center')
        self.subtotal_style.border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Estilo para datos normales
        self.data_style = NamedStyle(name='data')
        self.data_style.font = Font(size=10)
        self.data_style.alignment = Alignment(horizontal='left', vertical='center')
        self.data_style.border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Estilo para números
        self.number_style = NamedStyle(name='number')
        self.number_style.font = Font(size=10)
        self.number_style.alignment = Alignment(horizontal='center', vertical='center')
        self.number_style.border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Registrar estilos
        for style in [self.header_style, self.total_style, self.subtotal_style, 
                      self.data_style, self.number_style]:
            if style.name not in self.workbook.named_styles:
                self.workbook.add_named_style(style)
    
    def _apply_header_style(self, ws, row: int, start_col: int, end_col: int):
        """Aplica estilo de encabezado a una fila."""
        for col in range(start_col, end_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = Font(bold=True, color='FFFFFF', size=11)
            cell.fill = PatternFill(start_color='FF0000', end_color='FF0000', fill_type='solid')
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
    
    def _apply_total_style(self, ws, row: int, start_col: int, end_col: int):
        """Aplica estilo de total a una fila."""
        for col in range(start_col, end_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = Font(bold=True, color='000000', size=11)
            cell.fill = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
    
    def _apply_data_style(self, ws, row: int, start_col: int, end_col: int, alternate: bool = False):
        """Aplica estilo de datos a una fila."""
        fill_color = 'F0F0F0' if alternate else 'FFFFFF'
        for col in range(start_col, end_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = Font(size=10)
            cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type='solid')
            cell.alignment = Alignment(horizontal='center' if col > 1 else 'left', vertical='center')
            cell.border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
    
    def _write_dataframe(
        self,
        ws,
        df: pd.DataFrame,
        start_row: int = 1,
        start_col: int = 1,
        include_header: bool = True
    ) -> int:
        """
        Escribe un DataFrame en la hoja con estilos.
        
        Returns:
            Fila final después de escribir.
        """
        current_row = start_row
        n_cols = len(df.columns)
        
        # Encabezados
        if include_header:
            for col_idx, col_name in enumerate(df.columns, start=start_col):
                ws.cell(row=current_row, column=col_idx, value=col_name)
            self._apply_header_style(ws, current_row, start_col, start_col + n_cols - 1)
            current_row += 1
        
        # Datos
        for row_idx, row in df.iterrows():
            for col_idx, value in enumerate(row, start=start_col):
                ws.cell(row=current_row, column=col_idx, value=value)
            
            # Detectar si es fila de total
            first_val = str(row.iloc[0]).upper() if len(row) > 0 else ""
            if 'TOTAL' in first_val:
                self._apply_total_style(ws, current_row, start_col, start_col + n_cols - 1)
            elif 'SUBTOTAL' in first_val:
                for col in range(start_col, start_col + n_cols):
                    cell = ws.cell(row=current_row, column=col)
                    cell.font = Font(bold=True, color='FFFFFF', size=10)
                    cell.fill = PatternFill(start_color='0000FF', end_color='0000FF', fill_type='solid')
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                    cell.border = Border(
                        left=Side(style='thin'),
                        right=Side(style='thin'),
                        top=Side(style='thin'),
                        bottom=Side(style='thin')
                    )
            else:
                alternate = (current_row - start_row) % 2 == 0
                self._apply_data_style(ws, current_row, start_col, start_col + n_cols - 1, alternate)
            
            current_row += 1
        
        # Ajustar ancho de columnas
        from openpyxl.utils import get_column_letter
        for col_idx, col_name in enumerate(df.columns, start=start_col):
            max_length = len(str(col_name))
            for row in df[col_name]:
                max_length = max(max_length, len(str(row)))
            adjusted_width = min(max_length + 2, 50)
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = adjusted_width
        
        return current_row
    
    def _add_title(self, ws, title: str, row: int = 1, col: int = 1, colspan: int = 5):
        """Agrega un título a la hoja."""
        cell = ws.cell(row=row, column=col, value=title)
        cell.font = Font(bold=True, size=14, color='CC0000')
        cell.alignment = Alignment(horizontal='center')
        
        # Merge cells si hay más de una columna
        if colspan > 1:
            ws.merge_cells(
                start_row=row,
                start_column=col,
                end_row=row,
                end_column=col + colspan - 1
            )
    
    def _insert_image(self, ws, image_bytes: bytes, cell: str):
        """Inserta una imagen en la hoja."""
        try:
            img = XLImage(io.BytesIO(image_bytes))
            img.width = 600
            img.height = 400
            ws.add_image(img, cell)
        except Exception as e:
            print(f"Error insertando imagen: {e}")
    
    # ═══════════════════════════════════════════════════════════════════════
    # HOJAS DEL REPORTE
    # ═══════════════════════════════════════════════════════════════════════
    
    def _create_resumen_sheet(self):
        """Crea la hoja de resumen."""
        ws = self.workbook.active
        ws.title = "Resumen"
        
        periodo = self.report.periodo_principal
        if not periodo:
            return
        
        # Título
        self._add_title(ws, self.report.titulo, 1, 1, 3)
        
        # Información general
        info = [
            ("Jurisdicción:", self.report.jurisdiccion or "No especificada"),
            ("Período:", periodo.rango_fechas),
            ("Fecha de generación:", date.today().strftime("%d/%m/%Y")),
            ("", ""),
            ("RESUMEN DE DATOS", ""),
            ("Total Hechos Delictuales:", periodo.total_hechos),
            ("Total Mencionados:", periodo.total_mencionados),
            ("Total Aprehendidos:", periodo.total_aprehendidos),
        ]
        
        row = 3
        for label, value in info:
            ws.cell(row=row, column=1, value=label).font = Font(bold=True)
            ws.cell(row=row, column=2, value=value)
            row += 1
        
        # Ajustar anchos
        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 40
    
    def _create_table_sheet(
        self,
        name: str,
        title: str,
        df: pd.DataFrame,
        chart_bytes: Optional[bytes] = None
    ):
        """Crea una hoja con tabla y opcionalmente gráfico."""
        ws = self.workbook.create_sheet(title=name[:31])  # Max 31 chars
        
        # Título
        n_cols = len(df.columns) if not df.empty else 3
        self._add_title(ws, title, 1, 1, n_cols)
        
        # Período
        if self.report.periodo_principal:
            ws.cell(row=2, column=1, value=self.report.periodo_principal.rango_fechas)
            ws.cell(row=2, column=1).font = Font(bold=True, color='0000FF')
        
        # Tabla
        if not df.empty:
            final_row = self._write_dataframe(ws, df, start_row=4)
            
            # Gráfico
            if chart_bytes:
                self._insert_image(ws, chart_bytes, f'A{final_row + 2}')
    
    def _create_cuadro_referencia_sheet(
        self,
        name: str,
        title: str,
        df: pd.DataFrame
    ):
        """
        Crea hoja de cuadro de referencia con símbolos coloreados.
        
        Esta versión especial aplica el color del símbolo a la celda.
        """
        ws = self.workbook.create_sheet(title=name[:31])
        
        # Título
        n_cols = len(df.columns) if not df.empty else 4
        self._add_title(ws, title, 1, 1, n_cols)
        
        # Período
        if self.report.periodo_principal:
            ws.cell(row=2, column=1, value=self.report.periodo_principal.rango_fechas)
            ws.cell(row=2, column=1).font = Font(bold=True, color='0000FF')
        
        if df.empty:
            return
        
        # Escribir encabezados (sin columna Color)
        current_row = 4
        headers = [col for col in df.columns if col != 'Color']
        for col_idx, col_name in enumerate(headers, start=1):
            ws.cell(row=current_row, column=col_idx, value=col_name)
        self._apply_header_style(ws, current_row, 1, len(headers))
        current_row += 1
        
        # Escribir datos con símbolos coloreados
        for _, row in df.iterrows():
            color_hex = row.get('Color', '')
            
            col_idx = 1
            for col_name in headers:
                cell = ws.cell(row=current_row, column=col_idx, value=row[col_name])
                
                # Aplicar estilo base
                cell.alignment = Alignment(horizontal='center' if col_idx == 1 else 'left', vertical='center')
                cell.border = Border(
                    left=Side(style='thin'),
                    right=Side(style='thin'),
                    top=Side(style='thin'),
                    bottom=Side(style='thin')
                )
                
                # Aplicar color al símbolo (primera columna)
                if col_idx == 1 and color_hex:
                    # Convertir color hex a formato Excel (sin #)
                    excel_color = color_hex.replace('#', '')
                    # Para colores muy claros (blanco, amarillo), usar fondo negro
                    if excel_color.upper() in ['FFFFFF', 'FFFF00']:
                        cell.fill = PatternFill(start_color='000000', end_color='000000', fill_type='solid')
                    cell.font = Font(bold=True, color=excel_color, size=14)
                
                col_idx += 1
            
            current_row += 1
        
        # Ajustar anchos de columna
        ws.column_dimensions['A'].width = 8   # Símbolo
        ws.column_dimensions['B'].width = 12  # Tipo
        ws.column_dimensions['C'].width = 45  # Descripción
        ws.column_dimensions['D'].width = 12  # Cantidad
    
    # ═══════════════════════════════════════════════════════════════════════
    # CUADRO COMPARATIVO MEJORADO
    # ═══════════════════════════════════════════════════════════════════════
    
    def _create_comparative_sheet_enhanced(
        self,
        df: pd.DataFrame,
        chart_bytes: Optional[bytes] = None
    ):
        """
        Crea hoja de cuadro comparativo con formato visual mejorado.
        
        Aplica colores según tendencia:
        - Azul: Cuando el delito bajó (bueno)
        - Rojo: Cuando el delito subió (malo)
        - Amarillo: Fila de suma total
        
        Args:
            df: DataFrame con columnas de datos y 'color_fila' para estilos.
            chart_bytes: Gráfico comparativo en bytes PNG.
        """
        ws = self.workbook.create_sheet(title="Comparativa")
        
        if df.empty:
            return
        
        # Título
        self._add_title(ws, "CUADRO COMPARATIVO ENTRE PERÍODOS", 1, 1, 5)
        
        # Obtener columnas visibles (sin 'color_fila' y 'tendencia')
        display_cols = [col for col in df.columns if col not in ['color_fila', 'tendencia']]
        n_cols = len(display_cols)
        
        # Escribir encabezados (fila 3)
        current_row = 3
        for col_idx, col_name in enumerate(display_cols, start=1):
            cell = ws.cell(row=current_row, column=col_idx, value=col_name)
            cell.font = Font(bold=True, color='FFFFFF', size=10)
            cell.fill = PatternFill(start_color='FF0000', end_color='FF0000', fill_type='solid')
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
        
        # Ajustar altura de fila de encabezado
        ws.row_dimensions[current_row].height = 40
        current_row += 1
        
        # Escribir datos con colores según tendencia
        for _, row in df.iterrows():
            color_fila = row.get('color_fila', '#FFFFFF') if 'color_fila' in row.index else '#FFFFFF'
            tendencia = row.get('tendencia', '') if 'tendencia' in row.index else ''
            
            # Convertir color hex a formato Excel (sin #)
            excel_color = color_fila.replace('#', '') if color_fila else 'FFFFFF'
            
            # Determinar si es fila de total
            is_total = tendencia == 'total' or 'TOTAL' in str(row.iloc[0]).upper()
            
            for col_idx, col_name in enumerate(display_cols, start=1):
                value = row[col_name]
                cell = ws.cell(row=current_row, column=col_idx, value=value)
                
                # Estilo de borde común
                cell.border = Border(
                    left=Side(style='thin'),
                    right=Side(style='thin'),
                    top=Side(style='thin'),
                    bottom=Side(style='thin')
                )
                
                # Alineación
                cell.alignment = Alignment(
                    horizontal='center' if col_idx > 1 else 'left',
                    vertical='center'
                )
                
                # Aplicar colores según tipo de fila
                if is_total:
                    # Fila total: fondo amarillo, texto negro bold
                    cell.fill = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')
                    cell.font = Font(bold=True, color='000000', size=11)
                elif excel_color == 'FF0000':
                    # Rojo: delito subió (malo)
                    cell.fill = PatternFill(start_color='FF0000', end_color='FF0000', fill_type='solid')
                    cell.font = Font(color='FFFFFF', size=10)
                elif excel_color == '0000FF':
                    # Azul: delito bajó (bueno)
                    cell.fill = PatternFill(start_color='0000FF', end_color='0000FF', fill_type='solid')
                    cell.font = Font(color='FFFFFF', size=10)
                else:
                    # Sin cambio o normal
                    cell.fill = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')
                    cell.font = Font(color='000000', size=10)
            
            current_row += 1
        
        # Ajustar ancho de columnas
        from openpyxl.utils import get_column_letter
        column_widths = {
            1: 40,   # DELITO
            2: 18,   # Período 1
            3: 18,   # Período 2
            4: 35,   # Diferencia en cantidades
            5: 35,   # Diferencia porcentual
        }
        for col_idx, width in column_widths.items():
            if col_idx <= n_cols:
                col_letter = get_column_letter(col_idx)
                ws.column_dimensions[col_letter].width = width
        
        # Insertar gráfico comparativo debajo de la tabla
        if chart_bytes:
            chart_row = current_row + 2
            self._insert_image(ws, chart_bytes, f'A{chart_row}')
    
    # ═══════════════════════════════════════════════════════════════════════
    # EXPORTACIÓN PRINCIPAL
    # ═══════════════════════════════════════════════════════════════════════
    
    def export(
        self,
        output_path: str,
        tables: Dict[str, pd.DataFrame],
        charts: Optional[Dict[str, bytes]] = None,
        options: Optional[Dict[str, Any]] = None
    ) -> bool:
        """
        Exporta el reporte completo a Excel.
        
        Args:
            output_path: Ruta del archivo de salida.
            tables: Diccionario de tablas generadas.
            charts: Diccionario de gráficos (bytes PNG).
            options: Opciones de exportación (incluir_graficos, incluir_cuadro_ref, etc.)
        
        Returns:
            True si se exportó correctamente.
        """
        charts = charts or {}
        options = options or {}
        
        # Opciones por defecto
        incluir_graficos = options.get('incluir_graficos', True)
        incluir_cuadro_ref = options.get('incluir_cuadro_ref', True)
        incluir_mencionados = options.get('incluir_mencionados', True)
        incluir_matrices = options.get('incluir_matrices', True)
        incluir_comparativos = options.get('incluir_comparativos', True)
        
        try:
            # Hoja de resumen
            self._create_resumen_sheet()
            
            # Cuadro de referencia (con símbolos coloreados)
            if incluir_cuadro_ref and 'cuadro_referencia' in tables:
                self._create_cuadro_referencia_sheet(
                    "Cuadro Referencia",
                    "CUADRO DE REFERENCIA",
                    tables['cuadro_referencia']
                )
            
            # Delitos con modalidades
            if 'delitos' in tables:
                self._create_table_sheet(
                    "Delitos",
                    "DELITOS CON MODALIDADES",
                    tables['delitos'],
                    charts.get('delitos') if incluir_graficos else None
                )
            
            # Días de la semana
            if 'dias_semana' in tables:
                self._create_table_sheet(
                    "Días Semana",
                    "DÍAS DE LA SEMANA EN QUE OCURRIERON LOS HECHOS",
                    tables['dias_semana'],
                    charts.get('dias_semana') if incluir_graficos else None
                )
            
            # Franja horaria
            if 'franja_horaria' in tables:
                self._create_table_sheet(
                    "Franja Horaria",
                    "FRANJA HORARIA EN QUE OCURRIERON LOS HECHOS",
                    tables['franja_horaria'],
                    charts.get('franja_horaria') if incluir_graficos else None
                )
            
            # Movilidad
            if 'movilidad' in tables:
                self._create_table_sheet(
                    "Movilidad",
                    "MEDIOS DE MOVILIDAD UTILIZADOS",
                    tables['movilidad'],
                    charts.get('movilidad') if incluir_graficos else None
                )
            
            # Armas
            if 'armas' in tables:
                self._create_table_sheet(
                    "Armas",
                    "MEDIOS O ARMAS UTILIZADAS EN ROBOS AGRAVADOS",
                    tables['armas'],
                    charts.get('armas') if incluir_graficos else None
                )
            
            # Ámbito
            if 'ambito' in tables:
                self._create_table_sheet(
                    "Ámbito",
                    "AMBITO DE OCURRENCIA DELICTUAL",
                    tables['ambito'],
                    charts.get('ambito') if incluir_graficos else None
                )
            
            # Matriz delito × día
            if incluir_matrices and 'matriz_delito_dia' in tables:
                self._create_table_sheet(
                    "Delitos x Día",
                    "DELITOS POR DÍAS DE LA SEMANA",
                    tables['matriz_delito_dia'],
                    charts.get('delito_dia') if incluir_graficos else None
                )
            
            # Matriz delito × franja
            if incluir_matrices and 'matriz_delito_franja' in tables:
                self._create_table_sheet(
                    "Delitos x Franja",
                    "DELITOS POR FRANJA HORARIA",
                    tables['matriz_delito_franja'],
                    charts.get('delito_franja') if incluir_graficos else None
                )
            
            # Mencionados
            if incluir_mencionados and 'mencionados' in tables:
                self._create_table_sheet(
                    "Mencionados",
                    "MENCIONADOS COMO POSIBLES AUTORES MATERIALES",
                    tables['mencionados']
                )
            
            # Aprehendidos
            if 'aprehendidos' in tables:
                self._create_table_sheet(
                    "Aprehendidos",
                    "PERSONAS APREHENDIDAS",
                    tables['aprehendidos']
                )
            
            # Clasificación aprehendidos
            if 'aprehendidos_clasificacion' in tables:
                self._create_table_sheet(
                    "Clasif. Aprehendidos",
                    "CLASIFICACIÓN DE APREHENDIDOS",
                    tables['aprehendidos_clasificacion']
                )
            
            # Esclarecimiento
            if 'esclarecimiento' in tables:
                self._create_table_sheet(
                    "Esclarecimiento",
                    "ÍNDICE DE ESCLARECIMIENTO",
                    tables['esclarecimiento']
                )
            
            # Comparativa general - usar versión mejorada si hay datos detallados
            if incluir_comparativos:
                if 'comparativa_detallada' in tables and not tables['comparativa_detallada'].empty:
                    # Usar el nuevo método con colores y gráfico
                    self._create_comparative_sheet_enhanced(
                        tables['comparativa_detallada'],
                        charts.get('comparativa_delitos') if incluir_graficos else None
                    )
                elif 'comparativa_general' in tables:
                    self._create_table_sheet(
                        "Comparativa",
                        "CUADRO COMPARATIVO ENTRE PERÍODOS",
                        tables['comparativa_general']
                    )
            
            # Guardar
            Path(output_path).parent.mkdir(parents=True, exist_ok=True)
            
            # Intentar guardar con manejo de archivo bloqueado
            try:
                self.workbook.save(output_path)
            except PermissionError as pe:
                # El archivo puede estar abierto en Excel u otra app
                # Intentar con nombre alternativo
                from datetime import datetime
                base_path = Path(output_path)
                alt_name = f"{base_path.stem}_{datetime.now().strftime('%H%M%S')}{base_path.suffix}"
                alt_path = base_path.parent / alt_name
                try:
                    self.workbook.save(str(alt_path))
                    self._last_error = f"Archivo guardado como: {alt_name} (el original estaba bloqueado)"
                    return True
                except Exception:
                    raise PermissionError(
                        f"No se puede guardar el archivo. "
                        f"Por favor cierre el archivo '{base_path.name}' si está abierto en Excel "
                        f"o elija otra ubicación."
                    ) from pe
            
            return True
        
        except PermissionError as e:
            self._last_error = str(e)
            print(f"Error de permisos: {e}")
            return False
        
        except Exception as e:
            print(f"Error exportando a Excel: {e}")
            self._last_error = str(e)
            return False
